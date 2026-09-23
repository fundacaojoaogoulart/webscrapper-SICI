#!/usr/bin/env python3
"""Etapa A — triagem determinística de impacto documental.

Compara o estado do repositório com uma referência Git e classifica cada
arquivo alterado em: skip | regenerate | review_ai | manual_review.

Nenhuma IA é acionada aqui. Saída:

  * JSON impresso em stdout;
  * ``.docs-impact/impact.json`` (sempre);
  * ``.docs-impact/context.md`` (apenas quando há possível impacto).

Uso::

  python scripts/docs-impact.py                       # working tree vs HEAD
  python scripts/docs-impact.py --staged              # somente o que está em stage
  python scripts/docs-impact.py --base origin/main --head HEAD
  python scripts/docs-impact.py --test                # cenários sintéticos
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass

SCRIPT_DIR = Path(__file__).resolve().parent
MAPPING_FILE = SCRIPT_DIR / "docs-dependencies.json"

DIFF_CAP_TOTAL = 20000
DIFF_CAP_PER_FILE = 4000

CONVENTIONS = (
    "- Modifique somente os trechos afetados pela mudança.\n"
    "- Preserve títulos, terminologia, diagramas Mermaid e a organização existente.\n"
    "- Não reescreva arquivos inteiros quando um patch localizado for suficiente.\n"
    "- Não remova informações válidas nem substitua regras internas por interpretações próprias.\n"
    "- Não altere informações marcadas como \"externo não verificado\".\n"
    "- Registre dúvidas e conflitos para revisão humana, sem inventar fatos.\n"
    "- Se a documentação continuar correta, NÃO produza alterações.\n"
    "- Atualize `docs/data-lineage.md` (fluxo/rastreabilidade) e `docs/architecture.md`\n"
    "  (componentes/relações) quando a mudança afetar ambos.\n"
    "- Não regenere `docs/technical-review.md` (registro da auditoria inicial).\n"
)

SOURCE_EXTENSIONS = {
    ".py", ".txt", ".json", ".md", ".yml", ".yaml", ".toml",
    ".csv", ".xlsx", ".xls", ".pkl", ".joblib", ".sql", ".ini", ".cfg",
}


# --------------------------------------------------------------------------- #
# Git
# --------------------------------------------------------------------------- #
def run_git(args, repo):
    return subprocess.run(
        ["git", "-C", str(repo)] + args,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
    )


def repo_root():
    res = run_git(["rev-parse", "--show-toplevel"], SCRIPT_DIR.parent)
    out = res.stdout.strip()
    return Path(out) if out else SCRIPT_DIR.parent


def _parse_name_status(text: str):
    changes = []
    for line in text.splitlines():
        if not line:
            continue
        parts = line.split("\t")
        status_field = parts[0]
        code = status_field[0].upper()
        if code in ("R", "C") and len(parts) >= 3:
            changes.append({"status": code, "path": parts[2], "old_path": parts[1]})
        elif len(parts) >= 2:
            changes.append({"status": code, "path": parts[1], "old_path": None})
    return changes


def git_changed(repo, baseline, head, staged):
    if staged:
        res = run_git(["diff", "--cached", "--name-status", "-M"], repo)
        return _parse_name_status(res.stdout)
    if baseline and head:
        res = run_git(["diff", "--name-status", "-M", baseline, head], repo)
        return _parse_name_status(res.stdout)

    # Padrão: working tree vs HEAD (rastreados) + não rastreados não ignorados.
    res = run_git(["diff", "--name-status", "-M", "HEAD"], repo)
    changes = _parse_name_status(res.stdout)
    untracked = run_git(["ls-files", "--others", "--exclude-standard"], repo)
    for path in untracked.stdout.splitlines():
        if path:
            changes.append({"status": "A", "path": path, "old_path": None})
    return changes


# --------------------------------------------------------------------------- #
# Mapping + glob
# --------------------------------------------------------------------------- #
def load_dependencies(repo: Path):
    mapping_path = repo / "scripts" / "docs-dependencies.json"
    if not mapping_path.exists():
        raise SystemExit(f"mapeamento ausente: {mapping_path}")
    with open(mapping_path, encoding="utf-8") as fh:
        return json.load(fh)


def _glob_to_regex(pat: str) -> str:
    pat = pat.replace("\\", "/")
    out = []
    i, n = 0, len(pat)
    while i < n:
        c = pat[i]
        if c == "*":
            if i + 1 < n and pat[i + 1] == "*":
                out.append(".*")
                i += 2
                if i < n and pat[i] == "/":
                    out.append("(?:.*/)?")
                    i += 1
            else:
                out.append("[^/]*")
                i += 1
        elif c == "?":
            out.append("[^/]")
            i += 1
        elif c == "/":
            out.append("/")
            i += 1
        else:
            out.append(re.escape(c))
            i += 1
    return "^" + "".join(out) + "$"


def matches(pattern: str, path: str) -> bool:
    pattern = pattern.replace("\\", "/")
    path = path.replace("\\", "/")
    if "/" in pattern:
        return bool(re.match(_glob_to_regex(pattern), path))
    base = path.rsplit("/", 1)[-1]
    return bool(re.match(_glob_to_regex(pattern), base))


# --------------------------------------------------------------------------- #
# Classificação
# --------------------------------------------------------------------------- #
def classify(path: str, dep: dict):
    norm = path.replace("\\", "/")

    for pat in dep.get("regenerate", []):
        if matches(pat, norm):
            return {"classification": "regenerate", "reason": "apresentação/tooling: reconstruir o site"}

    for pat in dep.get("ignore", []):
        if matches(pat, norm):
            return {"classification": "skip", "reason": "arquivo ignorado"}

    for pat in dep.get("data_artifacts", []):
        if matches(pat, norm):
            return {"classification": "skip", "reason": "dados/registros (não altera a documentação)"}

    for entry in dep.get("files", []):
        if matches(entry["path"], norm):
            return {
                "classification": entry.get("classification", "review_ai"),
                "categories": entry.get("categories", []),
                "docs": entry.get("docs", []),
                "reason": "módulo mapeado",
            }

    ext = Path(norm).suffix.lower()
    if ext in SOURCE_EXTENSIONS:
        return {
            "classification": "manual_review",
            "categories": ["architecture_components"],
            "docs": [],
            "reason": "módulo novo/desconhecido (revisão humana)",
        }
    return {"classification": "skip", "reason": "sem impacto conhecido"}


def _looks_source_like(path: str) -> bool:
    return Path(path).suffix.lower() in SOURCE_EXTENSIONS


# --------------------------------------------------------------------------- #
# Diff relevante para o pacote de contexto
# --------------------------------------------------------------------------- #
def _schema_summary(repo: Path, path: str) -> str:
    """Resumo leve de schema de um .xlsx (abas + cabeçalhos), sem ler os dados."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "(openpyxl indisponível — não foi possível resumir o schema)"
    target = repo / path
    if not target.exists():
        return "(arquivo não encontrado no working tree)"
    try:
        wb = load_workbook(str(target), read_only=True, data_only=True)
        lines = [f"Abas: {', '.join(wb.sheetnames)}"]
        for ws in wb.worksheets:
            header = []
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                header = [str(c) for c in row if c is not None]
                break
            lines.append(f"  [{ws.title}] cabeçalho: {header}")
        wb.close()
        return "\n".join(lines)
    except Exception as exc:  # noqa: BLE001
        return f"(não foi possível abrir {path}: {exc})"


def _diff_for(repo: Path, change: dict, baseline: str, head: str) -> str:
    path = change["path"]
    ext = Path(path).suffix.lower()

    if ext in (".xlsx", ".xls"):
        return f"[schema de {path}]\n" + _schema_summary(repo, path)
    if ext in (".pkl", ".joblib", ".ico", ".png", ".jpg", ".gif", ".zip"):
        return f"[binário] {path} — diff não textual; revisão humana necessária"

    if change["status"] == "A" and not change.get("old_path"):
        target = repo / path
        if target.exists():
            try:
                return target.read_text(encoding="utf-8", errors="replace")[:DIFF_CAP_PER_FILE]
            except Exception:  # noqa: BLE001
                return f"[binário/não-textual] {path}"
        return f"(arquivo novo não encontrado: {path})"

    args = ["diff", "--no-color", "--"]
    if baseline and head:
        args = ["diff", "--no-color", baseline, head, "--", path]
    else:
        args = ["diff", "--no-color", "HEAD", "--", path]
    res = run_git(args, repo)
    return res.stdout or "(sem diff disponível)"


# --------------------------------------------------------------------------- #
# Triagem
# --------------------------------------------------------------------------- #
def triage(repo: Path, baseline: str | None, head: str | None, staged: bool) -> dict:
    dep = load_dependencies(repo)
    raw = git_changed(repo, baseline, head, staged)

    doc_by_id = {d["id"]: d for d in dep.get("docs", [])}

    changes = []
    for ch in raw:
        info = classify(ch["path"], dep)
        # Para renomeações, o destino manda; se for desconhecido, avalia o antigo.
        if info["classification"] == "manual_review" and info["reason"].startswith("módulo novo") and ch.get("old_path"):
            alt = classify(ch["old_path"], dep)
            if alt["classification"] != "manual_review":
                info = alt
                info["reason"] = "renomeação de módulo mapeado"
        changes.append({
            "path": ch["path"],
            "old_path": ch.get("old_path"),
            "status": ch["status"],
            "classification": info["classification"],
            "categories": info.get("categories", []),
            "docs": info.get("docs", []),
            "reason": info.get("reason", ""),
        })

    summary = {"skip": 0, "regenerate": 0, "review_ai": 0, "manual_review": 0}
    for c in changes:
        summary[c["classification"]] = summary.get(c["classification"], 0) + 1

    needs_agent = any(c["classification"] in ("review_ai", "manual_review") for c in changes)
    requires_human = any(c["classification"] == "manual_review" for c in changes)

    impact = {
        "baseline": baseline or "HEAD",
        "head": head or "worktree",
        "staged": staged,
        "needsAgent": needs_agent,
        "requiresHumanReview": requires_human,
        "summary": summary,
        "changes": changes,
    }

    candidate_ids = []
    for c in changes:
        for d in c["docs"]:
            if d not in candidate_ids:
                candidate_ids.append(d)
    candidate_paths = [doc_by_id[i]["path"] for i in candidate_ids if i in doc_by_id]
    impact["candidateDocs"] = candidate_paths

    candidate_hashes = {}
    for p in candidate_paths:
        fp = repo / p
        if fp.exists():
            candidate_hashes[p] = hashlib.sha256(fp.read_bytes()).hexdigest()
    impact["candidateHashes"] = candidate_hashes

    return impact


# --------------------------------------------------------------------------- #
# Pacote de contexto
# --------------------------------------------------------------------------- #
def build_context(repo: Path, impact: dict) -> str:
    dep = load_dependencies(repo)
    doc_by_id = {d["id"]: d for d in dep.get("docs", [])}

    relevant = [c for c in impact["changes"] if c["classification"] in ("review_ai", "manual_review")]

    lines = ["# Contexto de atualização da documentação", ""]
    lines.append("## Resumo")
    lines.append(f"- `needsAgent`: {impact['needsAgent']}")
    lines.append(f"- `requiresHumanReview`: {impact['requiresHumanReview']}")
    lines.append(f"- Arquivos relevantes: {', '.join(c['path'] for c in relevant) or '(nenhum)'}")
    if impact.get("candidateDocs"):
        lines.append("- Documentos candidatos: " + ", ".join(impact["candidateDocs"]))
    lines.append("")
    lines.append("## Convenções de atualização")
    lines.append("")
    lines.append(CONVENTIONS)
    lines.append("")

    lines.append("## Mudanças relevantes (diff)")
    total = 0
    for c in relevant:
        diff = _diff_for(repo, c, impact.get("baseline"), impact.get("head"))
        remaining = DIFF_CAP_TOTAL - total
        if remaining <= 0:
            lines.append(f"\n> (diff de {c['path']} omitido por limite de tamanho)")
            continue
        diff = diff[:min(len(diff), remaining, DIFF_CAP_PER_FILE)]
        total += len(diff)
        lines.append(f"\n### {c['status']} {c['path']}")
        if c.get("old_path"):
            lines.append(f"(renomeado de `{c['old_path']}`)")
        lines.append(f"classificação: {c['classification']} — {c['reason']}")
        lines.append(f"categorias: {', '.join(c['categories']) or '(nenhuma)'}")
        lines.append("```diff")
        lines.append(diff.rstrip("\n"))
        lines.append("```")
    lines.append("")

    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Escrita
# --------------------------------------------------------------------------- #
def write_outputs(repo: Path, impact: dict) -> None:
    impact_dir = repo / ".docs-impact"
    impact_dir.mkdir(parents=True, exist_ok=True)
    with open(impact_dir / "impact.json", "w", encoding="utf-8") as fh:
        json.dump(impact, fh, ensure_ascii=False, indent=2)

    if impact["needsAgent"]:
        context = build_context(repo, impact)
        with open(impact_dir / "context.md", "w", encoding="utf-8") as fh:
            fh.write(context)


# --------------------------------------------------------------------------- #
# Testes sintéticos
# --------------------------------------------------------------------------- #
def self_test() -> int:
    import shutil
    import tempfile

    real_dep = load_dependencies(SCRIPT_DIR.parent)
    failures = 0

    def check(name, cond):
        nonlocal failures
        if cond:
            print(f"  PASS  {name}")
        else:
            print(f"  FAIL  {name}")
            failures += 1

    with tempfile.TemporaryDirectory() as tmp:
        repo = Path(tmp)
        (repo / "scripts").mkdir(parents=True)
        (repo / "scripts" / "docs-dependencies.json").write_text(
            json.dumps(real_dep, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        # Arquivos de base (conteúdo mínimo).
        (repo / "scraper_sici_nome.py").write_text("seletor = '#x'\n", encoding="utf-8")
        (repo / "config.txt").write_text("[MACRO_AREAS]\ncasacivil = Gestão\n", encoding="utf-8")
        (repo / "gerar exe.txt").write_text("pyinstaller ...\n", encoding="utf-8")
        (repo / "MFE_Base.xlsx").write_bytes(b"PK\x03\x04")
        (repo / "model_fjg.pkl").write_bytes(b"\x80\x04model")
        (repo / "MFE_Atualizada.xlsx").write_bytes(b"PK\x03\x04data")
        (repo / "app.log").write_text("ruído\n", encoding="utf-8")

        def git(*args):
            return run_git(list(args), repo)

        git("init", "-q")
        git("config", "user.email", "t@t")
        git("config", "user.name", "t")
        git("add", ".")
        git("commit", "-q", "-m", "baseline")

        def run_triage():
            return triage(repo, "HEAD", None, False)

        # 1. Sem impacto (arquivo ignorado).
        (repo / "app.log").write_text("mais ruído\n", encoding="utf-8")
        imp = run_triage()
        check("cenário 1 (sem impacto) → needsAgent false", imp["needsAgent"] is False)

        # 2. Seletor/campo do scraper.
        git("checkout", "--", "app.log")
        (repo / "scraper_sici_nome.py").write_text("seletor = '#novo'\n", encoding="utf-8")
        imp = run_triage()
        cls = {c["path"]: c for c in imp["changes"]}
        check("cenário 2 (seletor) → review_ai", cls["scraper_sici_nome.py"]["classification"] == "review_ai")
        check("cenário 2 → needsAgent true", imp["needsAgent"] is True)
        check("cenário 2 → requer revisão humana? não", imp["requiresHumanReview"] is False)

        # 3. Regra de correspondência (config.txt).
        git("checkout", "--", "scraper_sici_nome.py")
        (repo / "config.txt").write_text("[MACRO_AREAS]\nsedecon = Social\n", encoding="utf-8")
        imp = run_triage()
        cls = {c["path"]: c for c in imp["changes"]}
        check("cenário 3 (regra) → review_ai", cls["config.txt"]["classification"] == "review_ai")

        # 4. Schema de entrada (MFE_Base.xlsx).
        git("checkout", "--", "config.txt")
        (repo / "MFE_Base.xlsx").write_bytes(b"PK\x03\x04novo")
        imp = run_triage()
        cls = {c["path"]: c for c in imp["changes"]}
        check("cenário 4 (schema) → manual_review", cls["MFE_Base.xlsx"]["classification"] == "manual_review")
        check("cenário 4 → requiresHumanReview true", imp["requiresHumanReview"] is True)

        # 5. Geração/execução do .exe.
        git("checkout", "--", "MFE_Base.xlsx")
        (repo / "gerar exe.txt").write_text("pyinstaller --onedir ...\n", encoding="utf-8")
        imp = run_triage()
        cls = {c["path"]: c for c in imp["changes"]}
        check("cenário 5 (exe) → review_ai + arquitetura", "architecture" in cls["gerar exe.txt"]["docs"])

        # 6. Revisão humana (binário de modelo).
        git("checkout", "--", "gerar exe.txt")
        (repo / "model_fjg.pkl").write_bytes(b"\x80\x04novo")
        imp = run_triage()
        cls = {c["path"]: c for c in imp["changes"]}
        check("cenário 6 (binário) → manual_review", cls["model_fjg.pkl"]["classification"] == "manual_review")

        # 7. Dados externos (valores) sem mudança de funcionamento.
        git("checkout", "--", "model_fjg.pkl")
        (repo / "MFE_Atualizada.xlsx").write_bytes(b"PK\x03\x04novos-dados")
        imp = run_triage()
        cls = {c["path"]: c for c in imp["changes"]}
        check("cenário 7 (dados) → skip", cls["MFE_Atualizada.xlsx"]["classification"] == "skip")
        check("cenário 7 → needsAgent false", imp["needsAgent"] is False)

    print(f"\nSintéticos: {'OK' if failures == 0 else f'{failures} falha(s)'}")
    return 0 if failures == 0 else 1


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Triagem determinística de impacto documental.")
    parser.add_argument("--staged", action="store_true", help="compara somente o que está em stage")
    parser.add_argument("--base", help="referência Git de base (ex.: origin/main)")
    parser.add_argument("--head", help="referência Git de topo (ex.: HEAD)")
    parser.add_argument("--test", action="store_true", help="executa os cenários sintéticos")
    parser.add_argument("--no-write", action="store_true", help="não grava .docs-impact/ (útil em CI)")
    args = parser.parse_args(argv)

    if args.test:
        return self_test()

    repo = repo_root()
    impact = triage(repo, args.base, args.head, args.staged)

    if not args.no_write:
        write_outputs(repo, impact)

    # JSON impresso ASCII-safe (portátil para CI); o arquivo é gravado em UTF-8.
    print(json.dumps(impact, ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
